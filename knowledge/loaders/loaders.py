"""多格式文档解析。"""
from __future__ import annotations

import re
from pathlib import Path
from core import config
from core.types import DocumentBlock


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def _clean_cell(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip()


def _table_to_markdown(rows: list[list[str]]) -> str:
    rows = [[_clean_cell(c) for c in row] for row in rows if any(str(c).strip() for c in row)]
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * ncol) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _load_docx(path: Path) -> list[DocumentBlock]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(str(path))
    blocks = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style else "") or ""
            m = re.match(r"Heading\s*(\d)", style, re.I)
            if m or style.lower().startswith("heading"):
                blocks.append(DocumentBlock(text, "heading", int(m.group(1)) if m else 1))
            elif style.lower().startswith(("list", "bullet", "number")):
                blocks.append(DocumentBlock(text, "list"))
            else:
                blocks.append(DocumentBlock(text))
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            md = _table_to_markdown([[c.text for c in row.cells] for row in table.rows])
            if md:
                blocks.append(DocumentBlock(md, "table"))
    return blocks


def _load_pdf(path: Path) -> list[DocumentBlock]:
    from pypdf import PdfReader
    blocks = []
    for i, page in enumerate(PdfReader(str(path)).pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(DocumentBlock(text, "paragraph", 0, i))
    return blocks


def _load_pptx(path: Path) -> list[DocumentBlock]:
    from pptx import Presentation
    blocks = []
    for i, slide in enumerate( Presentation(str(path)).slides, 1):
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                md = _table_to_markdown([[c.text for c in row.cells] for row in shape.table.rows])
                if md:
                    blocks.append(DocumentBlock(md, "table", 0, i))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if text:
                    blocks.append(DocumentBlock(text, "paragraph", 0, i))
    return blocks


def load_file(path: Path | str) -> list[DocumentBlock]:
    path = Path(path)
    def _load_plain(p: Path) -> list[DocumentBlock]:
        text = _read_text(p).strip()
        return [DocumentBlock(text)] if text else []
    loaders = {".docx": _load_docx, ".pdf": _load_pdf, ".pptx": _load_pptx, ".txt": _load_plain, ".md": _load_plain}
    loader = loaders.get(path.suffix.lower())
    if not loader:
        return []
    try:
        return loader(path)
    except Exception as exc:
        raise RuntimeError(f"解析失败 {path.name}: {exc}") from exc


def iter_documents(path: Path | str):
    root = Path(path)
    if root.is_file():
        if not root.name.startswith(".") and root.suffix.lower() in config.SUPPORTED_EXT:
            yield root
        return
    for p in sorted(root.rglob("*")):
        if p.is_file() and not p.name.startswith(".") and p.suffix.lower() in config.SUPPORTED_EXT:
            yield p
